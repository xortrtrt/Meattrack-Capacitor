from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from io import BytesIO, StringIO
from pathlib import Path
from uuid import uuid4
import csv

import pytest
from openpyxl import load_workbook

from app import repositories
from app.business_time import business_today, business_now
from app.database import fetch_one, get_transaction_cursor
from app.report_documents import csv_document, excel_document
from app.repositories import activity_reports as reports
from app.repositories import product_returns
from test_inventory_postgres import strict_inventory_database  # noqa: F401
from test_inventory_postgres import _payment_submission
from test_product_lifecycle import lifecycle, complete_refund, refund_proof  # noqa: F401
from test_sales_report_refunds import verified

pytestmark = pytest.mark.postgres


def build(d):
    return reports.build_report(d["sales"], business_today(), business_today())


def submit(d, payload=None, key=None):
    return reports.submit_report(d["sales"], business_today(), business_today(), reports.fingerprint(payload or build(d)),
                                 key or str(uuid4()), "Period reviewed", {"returns": "Inspected returned packs"})


def test_personal_attribution_deduplication_and_original_order_dates(lifecycle):  # noqa: F811
    d = lifecycle
    # An assigned order with no personal Sales action is not reportable by that Sales leader.
    unhandled = repositories.create_order("reseller", d["product"], 1, account_id=d["reseller"])
    with get_transaction_cursor() as cur:
        cur.execute("INSERT INTO activity_logs(account_id,action,entity_type,entity_id) VALUES(%s,'fulfilled_reseller_order','orders',%s)", (d["sales"], d["order"]))
    payload = build(d)
    assert [r["order_id"] for r in payload["orders"]] == [d["order"]]
    assert payload["summary"]["gross_sales"] == 1000
    assert unhandled["order_id"] not in [r["order_id"] for r in payload["orders"]]
    # A fulfillment today does not pull an older original order into today's report.
    with get_transaction_cursor() as cur:
        cur.execute("UPDATE orders SET order_date=order_date-interval '1 month' WHERE order_id=%s", (d["order"],))
    assert build(d)["orders"] == []
    older = business_today() - timedelta(days=40)
    assert reports.build_report(d["sales"], older, business_today())["summary"]["gross_sales"] == 1000


def test_snapshot_freezes_late_refunds_names_notes_and_exports(lifecycle):  # noqa: F811
    d = lifecycle
    case = verified(d)
    before = build(d)
    assert len(before["returns"]) == 1
    assert before["refunds"][0]["status"] == "pending"
    report_id = submit(d, before)
    frozen = reports.saved_report(report_id, d["sales"])
    csv_before = csv_document(frozen, "refunds")
    complete_refund(d, case)
    with get_transaction_cursor() as cur:
        cur.execute("UPDATE accounts SET name='Changed Sales Name' WHERE account_id=%s", (d["sales"],))
    assert reports.saved_report(report_id, d["sales"]) == frozen
    assert csv_document(reports.saved_report(report_id, d["owner"], owner=True), "refunds") == csv_before
    current = build(d)
    assert current["summary"]["retained_sales"] == 800
    assert current["refunds"][0]["processor"] == "Supervisor"
    assert frozen["summary"]["retained_sales"] == 1000
    listed = repositories.list_sales_reports(report_source="team_leader", team_leader_account_id=d["sales"])[0]
    assert listed["total_sales"] == 1000 and listed["submitted_by"] == frozen["author"]["name"]
    assert repositories.count_sales_reports(report_source="team_leader", q=frozen["author"]["name"]) == 1
    assert repositories.list_sales_reports(report_source="team_leader", q=frozen["author"]["name"])[0]["sales_report_id"] == report_id
    assert repositories.owner_export_rows("reports")[0]["total_sales"] == 1000
    assert reports.saved_report(report_id, d["inventory"]) is None
    assert frozen["notes"] == "Period reviewed" and frozen["section_notes"]["returns"] == "Inspected returned packs"
    assert fetch_one("SELECT COUNT(*) n FROM notifications WHERE source_type='sales_reports' AND source_id=%s", (report_id,))["n"] == 1
    with pytest.raises(Exception, match="append-only|immutable|trace|protected"):
        with get_transaction_cursor() as cur:
            cur.execute("UPDATE sales_report_activity_snapshots SET payload='{}' WHERE sales_report_id=%s", (report_id,))


def test_stale_preview_and_parallel_duplicate_submission(lifecycle):  # noqa: F811
    d = lifecycle
    old = build(d)
    verified(d)
    with pytest.raises(reports.StalePreview):
        submit(d, old)
    current = build(d)
    key = str(uuid4())
    with ThreadPoolExecutor(max_workers=2) as executor:
        ids = list(executor.map(lambda _: submit(d, current, key), range(2)))
    assert ids[0] == ids[1]
    assert fetch_one("SELECT COUNT(*) n FROM sales_report_activity_snapshots")["n"] == 1


def test_csv_and_excel_match_snapshot_and_neutralize_formula_text(lifecycle):  # noqa: F811
    d = lifecycle
    complete_refund(d, verified(d))
    payload = reports.with_notes(build(d), '=HYPERLINK("invalid")', {"sales": "+unexpected formula"})
    payload["sales"][0]["customer"] = "  =1+1"
    workbook = load_workbook(BytesIO(excel_document(payload)))
    assert workbook.sheetnames == ["Summary", "Sales", "Orders", "Returns", "Refunds"]
    for section in reports.SECTIONS:
        csv_rows = list(csv.DictReader(StringIO(csv_document(payload, section).lstrip("\ufeff"))))
        sheet = workbook[section.title()]
        assert sheet.freeze_panes == "C8" and sheet.print_title_rows == "$1:$7"
        assert sheet.page_setup.fitToWidth == 1 and sheet.page_setup.orientation == "landscape"
        assert sheet["A5"].value.startswith("'=")
        for row in sheet:
            assert all(cell.data_type != "f" for cell in row)
        if section != "summary":
            assert len(csv_rows) == len(payload[section])
    assert workbook["Sales"]["H8"].value == 800  # Retained sales remains a numeric cell.
    assert workbook["Sales"]["B8"].value.startswith("'")
    output = Path("outputs/activity-reports")
    output.mkdir(parents=True, exist_ok=True)
    (output / "sample.xlsx").write_bytes(excel_document(payload))


def test_return_only_handling_does_not_claim_sale_and_partial_full_refunds(lifecycle):  # noqa: F811
    d = lifecycle
    # Legacy orders with no recorded Sales action cannot be claimed as personal sales.
    with get_transaction_cursor() as cur:
        cur.execute("DELETE FROM activity_logs WHERE entity_type='orders' AND entity_id=%s", (d["order"],))
    reviewer = d["sales"]
    from test_product_lifecycle import create_case, receive
    case = create_case(d, quantity=10)
    product_returns.review_product_return(case, reviewer, True, "Reviewed all returned product packs")
    receive(d, case, quantity=10, authorize=False)
    product_returns.authorize_return_settlement(case, reviewer, "Proceed after custody verification", str(uuid4()))
    complete_refund(d, case)
    payload = reports.build_report(reviewer, business_today(), business_today())
    assert payload["summary"]["gross_sales"] == 0
    assert len(payload["returns"]) == len(payload["refunds"]) == 1
    assert payload["summary"]["completed_refunds"] == 1000
    assert payload["summary"]["retained_sales"] == 0


def test_migration_is_additive_and_repeatable_on_fresh_baseline(lifecycle):  # noqa: F811
    legacy = repositories.add_sales_report("team_leader", "Sales", business_today(), business_today(), 1000, 1,
                                            "Existing legacy report", account_id=lifecycle["sales"])
    migration = Path("database/migrations/037_sales_activity_reports.sql").read_text()
    with get_transaction_cursor() as cur:
        cur.execute("DROP TABLE sales_report_activity_snapshots")
        cur.execute(migration)
        cur.execute(migration)
    assert fetch_one("SELECT notes FROM sales_reports WHERE sales_report_id=%s", (legacy["sales_report_id"],))["notes"] == "Existing legacy report"
    assert submit(lifecycle) > 0


def test_failed_and_completed_cancellation_refunds_do_not_reduce_sales(lifecycle):  # noqa: F811
    d = lifecycle
    order = repositories.create_order("reseller", d["product"], 1, account_id=d["reseller"])
    attachments, payment = _payment_submission("activity-cancellation")
    payment["payment_amount"] = 100
    repositories.submit_order_payment(d["reseller"], order["order_id"], attachments, payment)
    repositories.decide_order(order["order_id"], "reject", team_leader_account_id=d["sales"],
                              rejection_reason="Customer requested cancellation", payment_received=True)
    refund_id = fetch_one("SELECT order_refund_id FROM order_refunds WHERE order_id=%s", (order["order_id"],))["order_refund_id"]
    repositories.start_order_refund(refund_id, d["supervisor"], "sales-supervisor")
    repositories.fail_order_refund(refund_id, d["supervisor"], "sales-supervisor", "Transfer provider unavailable")
    failed = build(d)
    assert failed["summary"]["pending_refunds"] == 100
    assert failed["summary"]["retained_sales"] == 1000
    assert failed["refunds"][0]["status"] == "failed"
    repositories.start_order_refund(refund_id, d["supervisor"], "sales-supervisor")
    repositories.complete_order_refund(refund_id, d["supervisor"], "sales-supervisor", refund_method="Bank transfer",
        transaction_reference=f"CANCEL{uuid4().hex[:14]}", transferred_at=business_now(),
        internal_note="Cancellation transfer confirmed against ledger", proof=refund_proof())
    completed = build(d)
    assert completed["summary"]["cancellation_refunds"] == 100
    assert completed["summary"]["sales_refunds"] == 0
    assert completed["summary"]["retained_sales"] == 1000

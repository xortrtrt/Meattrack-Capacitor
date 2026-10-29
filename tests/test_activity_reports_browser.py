from contextlib import contextmanager
from io import BytesIO
import socket
import threading
import time
import uvicorn
from pathlib import Path
from urllib.parse import urlsplit

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright

from app import main, web_sessions
from app.business_time import business_now
from app.repositories import activity_reports as reports
from test_inventory_postgres import strict_inventory_database  # noqa: F401
from test_product_lifecycle import lifecycle, complete_refund  # noqa: F401
from test_sales_report_refunds import verified

pytestmark = [pytest.mark.postgres, pytest.mark.e2e]


@contextmanager
def live_application():
    listener = socket.socket()
    listener.bind(("127.0.0.1", 0))
    port = listener.getsockname()[1]
    server = uvicorn.Server(uvicorn.Config(main.app, log_level="error", lifespan="off"))
    thread = threading.Thread(target=server.run, kwargs={"sockets": [listener]}, daemon=True)
    thread.start()
    try:
        for _ in range(250):
            if server.started:
                break
            time.sleep(.02)
        assert server.started, "The isolated browser-test server did not start."
        yield f"http://127.0.0.1:{port}"
    finally:
        server.should_exit = True
        thread.join(timeout=10)
        listener.close()


def test_prepare_submit_view_export_print_and_owner_access(lifecycle, monkeypatch):  # noqa: F811
    d = lifecycle
    complete_refund(d, verified(d))
    session = dict(account_id=d["sales"], role_key="team-leader", team_leader_role="sales",
                   portal_expires_at=int(business_now().timestamp()) + 3600)
    monkeypatch.setattr(web_sessions, "_load_session", lambda token: dict(session))
    monkeypatch.setattr(web_sessions, "_persist_session", lambda token, data, expires_seconds: session.update(data))
    output = Path("outputs/activity-reports")
    output.mkdir(parents=True, exist_ok=True)
    with live_application() as base_url, TestClient(main.app, base_url=base_url) as client, sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page = browser.new_page(viewport={"width": 1440, "height": 1000})
        errors = []
        page.on("pageerror", lambda error: errors.append(str(error)))
        assert page.goto(f"{base_url}/portal/team-leader/reports", wait_until="networkidle").status == 200
        expect(page.get_by_role("heading", name="Sales activity reports", exact=True)).to_be_visible()
        expect(page.get_by_text("Cumulative financial totals", exact=True)).to_have_count(0)
        page.get_by_role("button", name="Generate preview").click()
        page.wait_for_url("**/reports?**preview=1")
        page.wait_for_load_state("networkidle")
        expect(page.locator('[data-report-metric="retained_sales"]')).to_have_text("PHP 800.00")
        expect(page.get_by_role("tab", name="Sales 1", exact=True)).to_have_attribute("aria-selected", "true")
        for width in (375, 768, 1100, 1440):
            page.set_viewport_size({"width": width, "height": 1000})
            assert page.evaluate("document.documentElement.scrollWidth <= innerWidth + 1")
            page.screenshot(path=str(output / f"workspace-{width}.png"), full_page=True)
        page.get_by_label("Sales notes").fill("Sales reviewed for the period")
        page.get_by_role("tab", name="Sales 1", exact=True).focus()
        page.keyboard.press("End")
        expect(page.get_by_role("tab", name="Refunds 1", exact=True)).to_have_attribute("aria-selected", "true")
        expect(page.locator('#ar-panel-refunds')).to_contain_text("Supervisor")
        page.get_by_label("Refunds notes").fill("Supervisor completed the transfer")
        page.get_by_label("Report notes", exact=False).last.fill("All four sections reviewed")
        with page.expect_download() as download_info:
            page.get_by_role("button", name="Excel", exact=True).click()
        draft_download = download_info.value
        draft_download.save_as(str(output / "preview.xlsx"))
        draft_book = load_workbook(output / "preview.xlsx")
        assert draft_book["Sales"]["A6"].value == "Sales reviewed for the period"
        page.get_by_role("button", name="Submit to Owner", exact=True).click()
        page.wait_for_url("**/portal/team-leader/reports/*")
        page.wait_for_load_state("networkidle")
        report_id = int(urlsplit(page.url).path.rsplit("/", 1)[-1])
        expect(page.get_by_role("heading", name="Sales Activity Report", exact=True)).to_be_visible()
        expect(page.locator("main")).to_contain_text("All four sections reviewed")
        for section in reports.SECTIONS:
            response = client.get(f"/portal/team-leader/reports/{report_id}/csv?section={section}")
            assert response.status_code == 200 and response.content.startswith(b"\xef\xbb\xbf")
        response = client.get(f"/portal/team-leader/reports/{report_id}/xlsx")
        assert response.status_code == 200
        workbook = load_workbook(BytesIO(response.content))
        assert workbook["Sales"]["H8"].value == 800
        assert workbook["Summary"]["A5"].value == "All four sections reviewed"
        page.emulate_media(media="print")
        page.pdf(path=str(output / "report.pdf"), prefer_css_page_size=True, print_background=True)
        page.screenshot(path=str(output / "print-preview.png"), full_page=True)
        page.emulate_media(media="screen")
        page.goto(f"{base_url}/portal/team-leader/reports?view=history", wait_until="networkidle")
        expect(page.locator(".ar-history-card")).to_have_count(1)
        expect(page.get_by_role("link", name="View report", exact=True)).to_be_visible()
        session.update(account_id=d["supervisor"], role_key="sales-supervisor", team_leader_role="sales-supervisor")
        assert client.get(f"/portal/sales-supervisor/reports/{report_id}").status_code == 404
        session.update(account_id=d["owner"], role_key="owner", team_leader_role=None)
        page.goto(f"{base_url}/portal/owner/reports", wait_until="networkidle")
        page.get_by_role("link", name=f"Activity report #{report_id}", exact=False).click()
        page.wait_for_load_state("networkidle")
        expect(page.locator("main")).to_contain_text("Supervisor completed the transfer")
        assert client.get(f"/portal/owner/reports/{report_id}/csv?section=invalid").status_code == 404
        assert client.get(f"/portal/reseller/reports/{report_id}/xlsx").status_code == 404
        session.update(account_id=d["inventory"], role_key="team-leader", team_leader_role="inventory")
        assert client.get(f"/portal/team-leader/reports/{report_id}/xlsx", follow_redirects=False).status_code == 303
        assert not errors, errors
        browser.close()
